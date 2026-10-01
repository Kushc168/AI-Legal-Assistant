"""Text extraction and clause-aware chunking.

Offsets: every chunk records [start_offset, end_offset) into the contract's full_text, which is
the page texts joined with PAGE_SEPARATOR. Page numbers are 1-based.
"""

from __future__ import annotations

import bisect
import re
from dataclasses import dataclass
from pathlib import Path

PAGE_SEPARATOR = "\n\n"
SYNTHETIC_PAGE_CHARS = 3000
MAX_CHUNK_CHARS = 1200

_HEADING_PATTERNS = [
    # "Section 9.3", "Clause 4", "Article 12.1"
    re.compile(r"^(?:section|clause|article)\s+(\d+(?:\.\d+)*)\b[.:)]?\s*(.*)$", re.IGNORECASE),
    # "9.3 Liability", "4. TERM AND RENEWAL", "12) Notices"
    re.compile(r"^(\d{1,3}(?:\.\d{1,3})*)[.)]?\s+([A-Z(\"“][^\n]*)$"),
    # "A.2 The deliverables ..." (schedule sub-clauses)
    re.compile(r"^([A-Z]\.\d{1,3}(?:\.\d{1,3})*)[.)]?\s+([A-Z(\"“][^\n]*)$"),
    # "IV. CONFIDENTIALITY"
    re.compile(r"^([IVXL]{1,6})\.\s+([A-Z][^\n]*)$"),
]
_SCHEDULE = re.compile(r"^(schedule|annex(?:ure)?|appendix|exhibit)\s+([A-Z0-9]{1,3})\b[\s:.\-–—]*(.*)$", re.IGNORECASE)
_SENTENCE_SPLIT = re.compile(r"(?<=[.;:])\s+(?=[A-Z(\"'])")


class ExtractionError(Exception):
    pass


@dataclass
class Chunk:
    idx: int
    text: str
    start_offset: int
    end_offset: int
    page_start: int
    page_end: int
    clause_ref: str | None
    heading: str | None


@dataclass
class Document:
    pages: list[str]
    full_text: str
    page_starts: list[int]

    def page_of(self, offset: int) -> int:
        return bisect.bisect_right(self.page_starts, offset)


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _paginate(text: str, size: int = SYNTHETIC_PAGE_CHARS) -> list[str]:
    """Split text without page information into pseudo-pages at paragraph boundaries."""
    if "\f" in text:
        return [p for p in (_normalise(p) for p in text.split("\f")) if p]
    paragraphs = _normalise(text).split("\n\n")
    pages, current = [], ""
    for para in paragraphs:
        if current and len(current) + len(para) > size:
            pages.append(current)
            current = para
        else:
            current = f"{current}\n\n{para}" if current else para
    if current:
        pages.append(current)
    return pages


def extract_pages(path: Path, file_type: str) -> list[str]:
    file_type = file_type.lower()
    if file_type == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(path))
            pages = [_normalise(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:  # pypdf raises many types for malformed files
            raise ExtractionError(f"Could not read PDF: {exc}") from exc
        if not any(pages):
            raise ExtractionError(
                "No extractable text found. The PDF may be scanned; run OCR before uploading."
            )
        return [p for p in pages] or [""]
    if file_type == "docx":
        from docx import Document as DocxDocument

        try:
            doc = DocxDocument(str(path))
        except Exception as exc:
            raise ExtractionError(f"Could not read DOCX: {exc}") from exc
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return _paginate(text)
    if file_type in {"txt", "md"}:
        return _paginate(path.read_text(encoding="utf-8", errors="replace"))
    raise ExtractionError(f"Unsupported file type: {file_type}")


def build_document(pages: list[str]) -> Document:
    starts, offset = [], 0
    for page in pages:
        starts.append(offset)
        offset += len(page) + len(PAGE_SEPARATOR)
    return Document(pages=pages, full_text=PAGE_SEPARATOR.join(pages), page_starts=starts)


def parse_heading(line: str) -> tuple[str, str] | None:
    """Return (clause_ref, title) if the line looks like a clause heading."""
    # Only the start of the line matters: a clause is often one long paragraph ("3.2 This Agreement ...").
    line = line.strip()[:200]
    if not line:
        return None
    schedule = _SCHEDULE.match(line)
    if schedule:
        return f"{schedule.group(1).title()} {schedule.group(2).upper()}", schedule.group(3).strip()[:80]
    for pattern in _HEADING_PATTERNS:
        m = pattern.match(line)
        if m:
            ref = m.group(1)
            first = ref.split(".")[0]
            if first.isdigit() and int(first) > 99:
                return None
            title = (m.group(2) or "").strip()
            # Titles are short; a long sentence after a number is body text, not a heading.
            title = title.split(". ")[0][:80] if title else ""
            return ref, title
    return None


def _paragraph_spans(text: str) -> list[tuple[int, int]]:
    spans = []
    for m in re.finditer(r"[^\n](?:.|\n(?!\n))*", text):
        spans.append((m.start(), m.end()))
    return spans


def _split_long(text: str, start: int, end: int) -> list[tuple[int, int]]:
    """Split [start, end) into sentence-aligned pieces of at most MAX_CHUNK_CHARS."""
    if end - start <= MAX_CHUNK_CHARS:
        return [(start, end)]
    segment = text[start:end]
    cuts = [0] + [m.end() for m in _SENTENCE_SPLIT.finditer(segment)] + [len(segment)]
    pieces, piece_start = [], 0
    for prev, cut in zip(cuts, cuts[1:]):
        if cut - piece_start > MAX_CHUNK_CHARS and prev > piece_start:
            pieces.append((start + piece_start, start + prev))
            piece_start = prev
    pieces.append((start + piece_start, end))
    # Hard split anything still too long (no sentence boundaries).
    out = []
    for s, e in pieces:
        while e - s > MAX_CHUNK_CHARS:
            out.append((s, s + MAX_CHUNK_CHARS))
            s += MAX_CHUNK_CHARS
        out.append((s, e))
    return out


def chunk_document(doc: Document) -> list[Chunk]:
    """Group paragraphs into clause-aligned chunks with page and clause metadata."""
    text = doc.full_text
    groups: list[dict] = []
    clause_ref: str | None = None
    heading: str | None = None
    pending_heading_start: int | None = None

    for start, end in _paragraph_spans(text):
        para = text[start:end]
        first_line = para.split("\n", 1)[0]
        parsed = parse_heading(first_line)
        if parsed:
            clause_ref, heading = parsed
            heading = heading or None
            # A heading on its own line (no sentence) is attached to the next paragraph.
            if "\n" not in para.strip() and len(para) < 100 and not para.rstrip().endswith("."):
                pending_heading_start = start
                continue
            groups.append({"start": start, "end": end, "ref": clause_ref, "heading": heading})
            pending_heading_start = None
            continue
        start_at = pending_heading_start if pending_heading_start is not None else start
        pending_heading_start = None
        last = groups[-1] if groups else None
        if (
            last
            and last["ref"] == clause_ref
            and end - last["start"] <= MAX_CHUNK_CHARS
            and start_at == start
        ):
            last["end"] = end
        else:
            groups.append({"start": start_at, "end": end, "ref": clause_ref, "heading": heading})

    if pending_heading_start is not None:
        groups.append(
            {"start": pending_heading_start, "end": len(text), "ref": clause_ref, "heading": heading}
        )

    chunks: list[Chunk] = []
    for group in groups:
        for s, e in _split_long(text, group["start"], group["end"]):
            piece = text[s:e]
            if not piece.strip():
                continue
            chunks.append(
                Chunk(
                    idx=len(chunks),
                    text=piece,
                    start_offset=s,
                    end_offset=e,
                    page_start=doc.page_of(s),
                    page_end=doc.page_of(max(s, e - 1)),
                    clause_ref=group["ref"],
                    heading=group["heading"],
                )
            )
    return chunks


CONTRACT_TYPE_HINTS = [
    ("nda", r"non-?disclosure agreement|confidentiality agreement|mutual nda"),
    ("employment", r"employment agreement|contract of employment|\bemployer\b.{0,200}\bemployee\b"),
    ("services", r"services agreement|master services|statement of work|service provider|consulting agreement|professional services"),
    ("license", r"licen[cs]e agreement|licensor|licensee|end user licen"),
    ("lease", r"lease agreement|landlord|tenant|lessor|lessee"),
    ("supply", r"supply agreement|purchase order|supplier shall supply|goods"),
    ("sales", r"sale and purchase|sales agreement|buyer|seller"),
    ("partnership", r"partnership agreement|joint venture"),
]

PERSONAL_DATA_HINTS = (
    r"personal (data|information)|personally identifiable|\bPII\b|data subjects?|employee (data|records)"
    r"|customer data|end[- ]user data|health (data|information)"
)


def guess_contract_type(text: str) -> str:
    head = text[:6000].lower()
    for ctype, pattern in CONTRACT_TYPE_HINTS:
        if re.search(pattern, head, re.IGNORECASE | re.DOTALL):
            return ctype
    return "other"


def guess_personal_data(text: str) -> bool:
    return bool(re.search(PERSONAL_DATA_HINTS, text, re.IGNORECASE))


if __name__ == "__main__":
    sample = (
        "MASTER SERVICES AGREEMENT\n\n1. DEFINITIONS\n\n1.1 \"Services\" means the services.\n\n"
        "9. LIABILITY\n\n9.3 The Supplier's liability shall not be limited.\n\fPage two text."
    )
    document = build_document(_paginate(sample))
    for c in chunk_document(document):
        print(c.idx, c.clause_ref, c.page_start, repr(c.text[:60]))
