"""End-to-end API tests in offline mode: upload -> index -> risk report -> chat -> compare -> export."""

import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend.main import app

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def wait_for(client, url, done, timeout=60):
    deadline = time.time() + timeout
    while time.time() < deadline:
        data = client.get(url).json()
        if done(data):
            return data
        time.sleep(0.2)
    raise AssertionError(f"Timed out waiting for {url}: {data}")


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def loaded(client):
    created = client.post("/api/samples/load").json()
    ids = {c["filename"]: c["id"] for c in created}
    for cid in ids.values():
        wait_for(client, f"/api/contracts/{cid}", lambda d: d["status"] in {"indexed", "failed"})
        wait_for(client, f"/api/contracts/{cid}/risk-report", lambda d: d.get("status") in {"complete", "partial", "failed"})
    return ids


def test_samples_index(client, loaded):
    for cid in loaded.values():
        detail = client.get(f"/api/contracts/{cid}").json()
        assert detail["status"] == "indexed", detail.get("error")
        assert detail["page_count"] >= 1 and detail["chunk_count"] > 5


def test_msa_v1_risk_report(client, loaded):
    report = client.get(f"/api/contracts/{loaded['Sample_MSA_v1.txt']}/risk-report").json()
    assert report["status"] == "complete"
    types = {f["risk_type_id"] for f in report["findings"]}
    assert "FIN-01" in types, "unlimited liability should be detected"
    assert "BUS-01" in types, "automatic renewal should be detected"
    assert "OPS-03" in types, "missing SLA should be detected"
    unlimited = next(f for f in report["findings"] if f["risk_type_id"] == "FIN-01")
    assert unlimited["severity"] == "Critical"
    assert unlimited["clause_ref"] == "9.3" and unlimited["page"] >= 1
    assert "unlimited" in unlimited["quote"].lower()
    # Every clause finding's quote is verbatim from the stored contract text.
    doc = client.get(f"/api/contracts/{loaded['Sample_MSA_v1.txt']}/document").json()
    full = "\n\n".join(p["text"] for p in doc["pages"])
    for f in report["findings"]:
        if f["evidence_type"] == "clause":
            assert f["quote"] in full
    assert 0 <= report["overall_score"] <= 100
    assert report["recommendations"], "recommendations expected"
    for rec in report["recommendations"]:
        assert rec["item_ids"]


def test_nda_absence_and_applicability(client, loaded):
    report = client.get(f"/api/contracts/{loaded['Sample_NDA.txt']}/risk-report").json()
    findings = {f["risk_type_id"]: f for f in report["findings"]}
    assert findings["LEG-02"]["evidence_type"] == "absence"
    assert findings["LEG-02"]["page"] is None
    assert findings["LEG-02"]["confidence"] <= 0.85
    assert "LEG-01" not in findings, "the NDA has confidentiality clauses"
    na = {n["risk_type_id"] for n in report["not_applicable"]}
    assert {"CMP-01", "OPS-03"} <= na


def test_dismiss_rescores(client, loaded):
    report = client.get(f"/api/contracts/{loaded['Sample_MSA_v1.txt']}/risk-report").json()
    critical = next(f for f in report["findings"] if f["severity"] == "Critical")
    client.patch(f"/api/risk-findings/{critical['id']}", json={"status": "dismissed", "reviewer_note": "test"})
    after = client.get(f"/api/risk-reports/{report['id']}").json()
    assert after["overall_score"] < report["overall_score"]
    client.patch(f"/api/risk-findings/{critical['id']}", json={"status": "open"})


def test_chat_cites_verified_quotes(client, loaded):
    conv = client.post("/api/conversations", json={"contract_ids": [loaded["Sample_MSA_v1.txt"]]}).json()
    reply = client.post(f"/api/conversations/{conv['id']}/messages", json={"content": "What is the notice period for non-renewal?"}).json()
    answer = reply["assistant"]
    assert answer["found"]
    assert answer["content"].strip(), "answer text must not be empty"
    assert any("60 days" in c["quote"] for c in answer["citations"])
    assert all(not c["quote"][0].isdigit() or c["quote"][:3].count(".") for c in answer["citations"]), \
        "quotes must not start in the middle of a clause number"
    miss = client.post(f"/api/conversations/{conv['id']}/messages", json={"content": "zebra giraffe photosynthesis"}).json()
    assert not miss["assistant"]["found"]


def test_compare_versions(client, loaded):
    started = client.post("/api/comparisons", json={
        "contract_a_id": loaded["Sample_MSA_v1.txt"], "contract_b_id": loaded["Sample_MSA_v2.txt"]}).json()
    cmp = wait_for(client, f"/api/comparisons/{started['id']}", lambda d: d["status"] in {"complete", "failed"})
    assert cmp["status"] == "complete", cmp.get("error")
    by_cat = {c["category_id"]: c for c in cmp["changes"]}
    assert len(by_cat) == 14
    notice = by_cat["notice_period"]
    assert notice["status"] == "modified"
    assert "60" in (notice["value_a"] or "") and "90" in (notice["value_b"] or "")
    assert by_cat["service_levels"]["status"] == "added"
    assert by_cat["confidentiality"]["status"] == "unchanged"
    assert cmp["risk_delta"] is not None
    assert cmp["anchors"], "alignment anchors for synchronised scrolling"
    md = client.get(f"/api/comparisons/{cmp['id']}/export?format=md")
    assert md.status_code == 200 and "Contract comparison summary" in md.text
    pdf = client.get(f"/api/comparisons/{cmp['id']}/export?format=pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert client.get(f"/api/comparisons/{cmp['id']}/export?format=docx").status_code == 501


def test_risk_export(client, loaded):
    report = client.get(f"/api/contracts/{loaded['Sample_MSA_v1.txt']}/risk-report").json()
    md = client.get(f"/api/risk-reports/{report['id']}/export?format=md")
    assert md.status_code == 200 and "Overall contract risk" in md.text
    pdf = client.get(f"/api/risk-reports/{report['id']}/export?format=pdf")
    assert pdf.content[:4] == b"%PDF"


def test_dashboard_and_analytics(client, loaded):
    dash = client.get("/api/dashboard").json()
    assert dash["total_contracts"] >= 3
    assert dash["portfolio_risk"]["score"] is not None
    assert dash["ai_conversations"] >= 1
    assert len(dash["recent_searches"]) >= 1
    assert client.get("/api/analytics").json()["contracts_analysed"] >= 3


def test_pdf_and_docx_uploads(client):
    from docx import Document
    from fpdf import FPDF

    pdf = FPDF()
    for page in range(2):
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 6, f"{page + 4}. TERMINATION\n\n{page + 4}.1 Either party may terminate this Agreement "
                             "on 120 days' prior written notice.")
    resp = client.post("/api/contracts", files={"file": ("test.pdf", bytes(pdf.output()), "application/pdf")})
    assert resp.status_code == 201
    detail = wait_for(client, f"/api/contracts/{resp.json()['id']}", lambda d: d["status"] in {"indexed", "failed"})
    assert detail["status"] == "indexed" and detail["page_count"] == 2

    import io

    doc = Document()
    doc.add_paragraph("1. PAYMENT")
    doc.add_paragraph("1.1 Invoices are payable within 30 days.")
    buf = io.BytesIO()
    doc.save(buf)
    resp = client.post("/api/contracts", files={"file": ("test.docx", buf.getvalue(), "application/octet-stream")})
    detail = wait_for(client, f"/api/contracts/{resp.json()['id']}", lambda d: d["status"] in {"indexed", "failed"})
    assert detail["status"] == "indexed"


def test_upload_validation(client):
    assert client.post("/api/contracts", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 400
    assert client.post("/api/contracts", files={"file": ("x.txt", b"", "text/plain")}).status_code == 400
